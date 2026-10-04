#!/usr/bin/env python3
"""Build an idempotent SQL import of public Primary and JSS schools for every state.

Usage: seed-national-schools.py [workbook.xlsx] > national-schools.sql
Then apply it with psql (ON_ERROR_STOP=1). Existing schools are never changed:
rows are inserted with ON CONFLICT DO NOTHING, so schools already referenced by
plans keep their ids and values. Yobe is skipped because it was seeded earlier.

Learner counts are deliberately synthetic (same stable formula as the Yobe
seed) so infrastructure models can be tested. Replace them with an
authoritative DNEMIS / Annual School Census import for real planning.
"""

import csv
import hashlib
import io
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

# state_code, name, lga, level, town, category, location, enrolment_male, enrolment_female
SchoolRow = tuple[str, str, str, str, str, str, str, int, int]

WORKBOOK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/Users/muhammad/Downloads/School list-1.xlsx")
SHEETS = {"PRIMARY": "Primary", "JSS": "JSS"}
SKIP_STATES = {"YO"}
STATE_CODES = {
    "ABIA": "AB", "ADAMAWA": "AD", "AKWA-IBOM": "AK", "ANAMBRA": "AN", "BAUCHI": "BA", "BAYELSA": "BY",
    "BENUE": "BE", "BORNO": "BO", "CROSS-RIVER": "CR", "DELTA": "DE", "EBONYI": "EB", "EDO": "ED",
    "EKITI": "EK", "ENUGU": "EN", "FCT-ABUJA": "FC", "GOMBE": "GO", "IMO": "IM", "JIGAWA": "JI",
    "KADUNA": "KD", "KANO": "KN", "KATSINA": "KT", "KEBBI": "KE", "KOGI": "KO", "KWARA": "KW",
    "LAGOS": "LA", "NASARAWA": "NA", "NIGER": "NI", "OGUN": "OG", "ONDO": "ON", "OSUN": "OS",
    "OYO": "OY", "PLATEAU": "PL", "RIVERS": "RI", "SOKOTO": "SO", "TARABA": "TA", "YOBE": "YO",
    "ZAMFARA": "ZA",
}


def clean(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def sample_enrolment(name: str, lga: str, level: str) -> tuple[int, int]:
    """Stable demo figures spanning the three infrastructure models (see seed-yobe-schools.py)."""
    digest = int(hashlib.sha256(f"{name}|{lga}|{level}".encode()).hexdigest()[:12], 16)
    bands = ((120, 240), (241, 320), (321, 560))
    minimum, maximum = bands[digest % len(bands)]
    total = minimum + ((digest >> 3) % (maximum - minimum + 1))
    female_share = 46 + ((digest >> 11) % 7)
    female = round(total * female_share / 100)
    return total - female, female


def read_rows(book: Workbook) -> list[SchoolRow]:
    rows: list[SchoolRow] = []
    seen: set[tuple[str, str, str, str]] = set()
    unknown: Counter[str] = Counter()
    for sheet_name, level in SHEETS.items():
        for row in book[sheet_name].iter_rows(min_row=3, values_only=True):
            state, lga, name, town, category, location = (clean(value) for value in row[:6])
            code = STATE_CODES.get(state.upper())
            if not code:
                if state:
                    unknown[state] += 1
                continue
            if code in SKIP_STATES or category.upper() != "PUBLIC" or not name or not lga:
                continue
            lga = lga.upper()
            key = (code, name.upper(), lga, level)
            if key in seen:
                continue
            seen.add(key)
            location = location.title() if location.title() in {"Rural", "Urban"} else "Urban"
            male, female = sample_enrolment(name, lga, level)
            rows.append((code, name, lga, level, town, "Public", location, male, female))
    if unknown:
        raise SystemExit(f"Unrecognised state names: {dict(unknown)}")
    return rows


def main() -> None:
    if not WORKBOOK.exists():
        raise SystemExit(f"School workbook not found: {WORKBOOK}")
    rows = read_rows(load_workbook(WORKBOOK, read_only=True, data_only=True))
    data = io.StringIO()
    csv.writer(data, lineterminator="\n").writerows(rows)
    sys.stdout.write(
        "BEGIN;\n"
        "CREATE TEMP TABLE school_import (state_code TEXT, name TEXT, lga TEXT, level TEXT, town TEXT, category TEXT,"
        " location TEXT, enrolment_male INTEGER, enrolment_female INTEGER) ON COMMIT DROP;\n"
        "COPY school_import FROM STDIN WITH (FORMAT csv);\n"
        + data.getvalue() + "\\.\n"
        "INSERT INTO schools (state_code, name, lga, level, town, category, location, enrolment_male, enrolment_female)\n"
        "SELECT state_code, name, lga, level, COALESCE(town, ''), category, location, enrolment_male, enrolment_female FROM school_import\n"
        "ON CONFLICT (state_code, name, lga, level) WHERE dnemis_id IS NULL DO NOTHING;\n"
        "COMMIT;\n"
    )
    per_state = Counter(row[0] for row in rows)
    print(f"Prepared {len(rows)} public Primary/JSS schools across {len(per_state)} states: "
          + ", ".join(f"{code}={count}" for code, count in sorted(per_state.items())), file=sys.stderr)


if __name__ == "__main__":
    main()
