#!/usr/bin/env python3
"""Load the Yobe rows from the supplied school workbook into local PostgreSQL."""

import csv
import hashlib
import io
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/Users/muhammad/Downloads/School list-1.xlsx")
SHEETS = {"PRIMARY": "Primary", "JSS": "JSS", "SSS": "SSS"}
LGA_FIXES = {"BARDE": "BADE", "BOSARI": "BURSARI", "TARMUA": "TARMUWA"}


def clean(value):
    return " ".join(str(value or "").strip().split())


def sample_enrolment(name, lga, level):
    """Return stable demo figures spanning all three infrastructure models.

    These values are deliberately synthetic and must be replaced by an
    authoritative DNEMIS/Annual School Census import for production planning.
    """
    digest = int(hashlib.sha256(f"{name}|{lga}|{level}".encode()).hexdigest()[:12], 16)
    bands = ((120, 240), (241, 320), (321, 560))
    minimum, maximum = bands[digest % len(bands)]
    total = minimum + ((digest >> 3) % (maximum - minimum + 1))
    female_share = 46 + ((digest >> 11) % 7)
    female = round(total * female_share / 100)
    return total - female, female


def main():
    if not WORKBOOK.exists():
        raise SystemExit(f"School workbook not found: {WORKBOOK}")

    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    rows = []
    seen = set()
    for sheet_name, level in SHEETS.items():
        sheet = book[sheet_name]
        for row in sheet.iter_rows(min_row=3, values_only=True):
            state, lga, name, town, category, location = (clean(value) for value in row[:6])
            if state.upper() != "YOBE" or not name or not lga:
                continue
            lga = LGA_FIXES.get(lga.upper(), lga.upper())
            location = location.title()
            location = location if location in {"Rural", "Urban"} else "Urban"
            key = (name.upper(), lga, level)
            if key not in seen:
                seen.add(key)
                male, female = sample_enrolment(name, lga, level)
                rows.append((name, lga, level, town, category.title(), location, male, female))

    data = io.StringIO()
    writer = csv.writer(data, lineterminator="\n")
    writer.writerows(rows)
    sql = """
ALTER TABLE schools ADD COLUMN IF NOT EXISTS town TEXT NOT NULL DEFAULT '';
ALTER TABLE schools ADD COLUMN IF NOT EXISTS category TEXT NOT NULL DEFAULT '';
ALTER TABLE schools DROP CONSTRAINT IF EXISTS schools_name_lga_key;
ALTER TABLE schools DROP CONSTRAINT IF EXISTS schools_name_lga_level_key;
ALTER TABLE schools ADD CONSTRAINT schools_name_lga_level_key UNIQUE (name, lga, level);
CREATE TABLE IF NOT EXISTS users (
  id SERIAL PRIMARY KEY, email TEXT NOT NULL UNIQUE, full_name TEXT NOT NULL,
  role TEXT NOT NULL CHECK (role IN ('Data Entry Officer', 'Reviewer', 'Executive Secretary')),
  password_hash TEXT NOT NULL, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
);
CREATE TABLE IF NOT EXISTS sessions (
  token TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  expires_at TIMESTAMPTZ NOT NULL, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
);
-- Provision demo accounts separately; the school import must not reset passwords.
TRUNCATE infrastructure_lines, schools RESTART IDENTITY CASCADE;
COPY schools (name, lga, level, town, category, location, enrolment_male, enrolment_female) FROM STDIN WITH (FORMAT csv);
""" + data.getvalue() + "\\.\n" + """
INSERT INTO action_plans (state_code, start_year, end_year) VALUES ('YO', 2025, 2025) ON CONFLICT DO NOTHING;
INSERT INTO infrastructure_lines (school_id, plan_id, code, project_type, quantity, rationale, strategy, longitude, latitude, duration, unit_cost)
SELECT id, (SELECT id FROM action_plans WHERE state_code = 'YO' AND start_year = 2025 AND end_year = 2025), 'UBC/SUBEB/NC/' || TO_CHAR(sequence, 'FM000') || '/2025', 'six-classrooms', 1,
  'Overcrowded classrooms', 'NCB', '', '',
  (SELECT duration FROM construction_types WHERE id = 'six-classrooms'),
  (SELECT unit_cost FROM construction_types WHERE id = 'six-classrooms')
FROM (
  SELECT id, ROW_NUMBER() OVER (ORDER BY lga, name) AS sequence
  FROM schools
  ORDER BY lga, name
  LIMIT 4
) AS initial_lines;
"""
    result = subprocess.run(
        ["docker", "compose", "exec", "-T", "postgres", "psql", "-v", "ON_ERROR_STOP=1", "-U", "ubec_app", "-d", "ubec"],
        input=sql, text=True, capture_output=True,
    )
    if result.returncode:
        raise SystemExit(result.stderr.strip() or "Unable to seed PostgreSQL.")
    model_counts = [0, 0, 0]
    for row in rows:
        total = row[-2] + row[-1]
        model_counts[0 if total <= 240 else 1 if total <= 320 else 2] += 1
    print(
        f"Seeded {len(rows)} Yobe schools with sample learner counts "
        f"(small={model_counts[0]}, medium={model_counts[1]}, large={model_counts[2]}) "
        "and 4 infrastructure project lines."
    )


if __name__ == "__main__":
    main()
